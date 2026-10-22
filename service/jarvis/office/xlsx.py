"""Edit .xlsx workbooks without losing anything Jarvis did not touch.

Saving a workbook through openpyxl drops charts, pivot tables, images and other parts it does not
model, so Jarvis never re-saves a user's workbook that way. Reading uses openpyxl (formulas, cached
values, shared formulas expanded); writing patches only the changed cells' XML inside the zip and
copies every other part byte for byte. Formula results are recomputed (formulas.py) and written as
cached values, and the workbook is flagged to recalculate fully when Excel or LibreOffice opens it.
"""

from __future__ import annotations

import html
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from jarvis.office import formulas as fx

MAX_CELLS = 400_000
MAX_UNZIPPED = 300 * 1024 * 1024
MAX_FILE = 50 * 1024 * 1024
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class XlsxError(ValueError):
    pass


@dataclass
class Sheet:
    name: str
    xml_path: str
    cells: dict[str, Any] = field(default_factory=dict)      # raw: numbers, text, '=formula'
    cached: dict[str, Any] = field(default_factory=dict)     # values as last computed by an app
    max_row: int = 0
    max_col: int = 0


@dataclass
class Book:
    path: Path
    sheets: dict[str, Sheet]

    def raw(self) -> dict[str, dict[str, Any]]:
        return {name: dict(s.cells) for name, s in self.sheets.items()}


def _check_zip(path: Path) -> zipfile.ZipFile:
    if path.stat().st_size > MAX_FILE:
        raise XlsxError(f"{path.name} is larger than {MAX_FILE // 1048576} MB")
    try:
        zf = zipfile.ZipFile(path)
    except zipfile.BadZipFile:
        raise XlsxError(f"{path.name} is not a valid .xlsx file") from None
    if sum(i.file_size for i in zf.infolist()) > MAX_UNZIPPED:
        raise XlsxError(f"{path.name} expands to more than {MAX_UNZIPPED // 1048576} MB")
    if any(i.filename.startswith("/") or ".." in i.filename.split("/") for i in zf.infolist()):
        raise XlsxError(f"{path.name} contains unsafe paths")
    return zf


def _sheet_paths(zf: zipfile.ZipFile) -> dict[str, str]:
    wb = zf.read("xl/workbook.xml").decode("utf-8")
    rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = {}
    for m in re.finditer(r"<Relationship\b([^>]*)/?>", rels):
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', m.group(1)))
        if "Id" in attrs and "Target" in attrs:
            t = attrs["Target"]
            targets[attrs["Id"]] = t.lstrip("/") if t.startswith("/") else "xl/" + t
    out = {}
    for m in re.finditer(r"<(?:\w+:)?sheet\b([^>]*)/?>", wb):
        attrs = dict(re.findall(r'([\w:]+)="([^"]*)"', m.group(1)))
        rid = next((v for k, v in attrs.items() if k.endswith(":id")), "")
        if "name" in attrs and rid in targets:
            out[html.unescape(attrs["name"])] = targets[rid]
    return out


def load(path: Path) -> Book:
    from openpyxl import load_workbook

    zf = _check_zip(path)
    paths = _sheet_paths(zf)
    zf.close()
    wb = load_workbook(path, data_only=False, read_only=False)
    wbv = load_workbook(path, data_only=True, read_only=False)
    sheets: dict[str, Sheet] = {}
    total = 0
    for ws in wb.worksheets:
        sheet = Sheet(ws.title, paths.get(ws.title, ""), max_row=ws.max_row or 0, max_col=ws.max_column or 0)
        wsv = wbv[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                total += 1
                if total > MAX_CELLS:
                    raise XlsxError("workbook too large to edit safely")
                value = cell.value
                if hasattr(value, "text"):          # ArrayFormula / DataTableFormula
                    value = "=" + str(getattr(value, "text", "")).lstrip("=")
                sheet.cells[cell.coordinate] = value if not hasattr(value, "isoformat") else value.isoformat()
                cached = wsv[cell.coordinate].value
                if cached is not None:
                    sheet.cached[cell.coordinate] = cached if not hasattr(cached, "isoformat") else cached.isoformat()
        sheets[ws.title] = sheet
    wb.close()
    wbv.close()
    return Book(path, sheets)


def outline(book: Book, rows: int = 12) -> list[dict]:
    """Per sheet: size, header row and the first rows as displayed values (for the model and the UI)."""
    out = []
    for s in book.sheets.values():
        grid = []
        for r in range(1, min(s.max_row, rows) + 1):
            line = []
            for c in range(1, min(s.max_col, 16) + 1):
                ref = f"{fx.col_letters(c)}{r}"
                raw = s.cells.get(ref)
                shown = s.cached.get(ref, raw) if isinstance(raw, str) and raw.startswith("=") else raw
                line.append(shown)
            grid.append(line)
        formulas = sum(1 for v in s.cells.values() if isinstance(v, str) and v.startswith("="))
        out.append({"sheet": s.name, "rows": s.max_row, "cols": s.max_col, "formulas": formulas, "preview": grid})
    return out


def _is_formula(v) -> bool:
    return isinstance(v, str) and v.startswith("=")


def find_rows(book: Book, query: str, limit: int = 20) -> list[dict]:
    """Rows whose text cells mention `query` (any word, case-insensitive), with their numbers."""
    words = [w for w in re.findall(r"\w+", query.casefold()) if len(w) > 1]
    if not words:
        return []
    hits = []
    for s in book.sheets.values():
        by_row: dict[int, dict[str, Any]] = {}
        for ref, v in s.cells.items():
            m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
            if m:
                by_row.setdefault(int(m.group(2)), {})[ref] = v
        for r, cells in sorted(by_row.items()):
            labels = [str(v) for v in cells.values() if isinstance(v, str) and not _is_formula(v)]
            text = " ".join(labels).casefold()
            score = sum(1 for w in words if w in text or any(t.startswith(w[:max(4, len(w) - 2)]) for t in text.split()))
            if not score:
                continue
            numbers = []
            for ref, v in sorted(cells.items(), key=lambda kv: fx.col_index(re.match(r"[A-Z]+", kv[0]).group())):
                if _is_formula(v):
                    numbers.append({"cell": ref, "formula": v, "value": s.cached.get(ref)})
                elif isinstance(v, (int, float)) and not isinstance(v, bool):
                    numbers.append({"cell": ref, "value": v})
            hits.append({"sheet": s.name, "row": r, "label": " | ".join(labels)[:160], "score": score, "cells": numbers})
    hits.sort(key=lambda h: (-h["score"], h["sheet"], h["row"]))
    return hits[:limit]


def recompute(book: Book, updates: dict[tuple[str, str], Any]) -> tuple[dict[tuple[str, str], Any], list[str]]:
    """All formula values after `updates`; and the formula cells that could not be evaluated."""
    raw = book.raw()
    for (sheet, ref), value in updates.items():
        raw.setdefault(sheet, {})[ref] = value
    ev = fx.Workbook(raw)
    values: dict[tuple[str, str], Any] = {}
    unsupported: list[str] = []
    for sheet, cells in raw.items():
        for ref, v in cells.items():
            if _is_formula(v):
                try:
                    values[(sheet, ref)] = ev.value(sheet, ref)
                except fx.Unsupported as exc:
                    unsupported.append(f"{sheet}!{ref} ({exc})")
    return values, unsupported


# ---------------------------------------------------------------- writing

_ROW = r'<row\b[^>]*\br="{r}"[^>]*?(?:/>|>.*?</row>)'
_CELL = r'<c\b[^>]*\br="{ref}"[^>]*?(?:/>|>.*?</c>)'


def _xml_text(value: str) -> str:
    return html.escape(value, quote=False)


def _cell_xml(ref: str, value: Any, style: str, cached: Any = None) -> str:
    s = f' s="{style}"' if style else ""
    if _is_formula(value):
        f = f"<f>{_xml_text(value[1:])}</f>"
        if cached is None or (isinstance(cached, str) and cached in fx.ERRORS):
            return f'<c r="{ref}"{s}>{f}</c>'
        if isinstance(cached, bool):
            return f'<c r="{ref}"{s} t="b">{f}<v>{int(cached)}</v></c>'
        if isinstance(cached, (int, float)):
            return f'<c r="{ref}"{s}>{f}<v>{_num(cached)}</v></c>'
        return f'<c r="{ref}"{s} t="str">{f}<v>{_xml_text(str(cached))}</v></c>'
    if isinstance(value, bool):
        return f'<c r="{ref}"{s} t="b"><v>{int(value)}</v></c>'
    if isinstance(value, (int, float)):
        return f'<c r="{ref}"{s}><v>{_num(value)}</v></c>'
    if value is None or value == "":
        return f'<c r="{ref}"{s}/>'
    text = str(value)
    space = ' xml:space="preserve"' if text != text.strip() else ""
    return f'<c r="{ref}"{s} t="inlineStr"><is><t{space}>{_xml_text(text)}</t></is></c>'


def _num(v: float) -> str:
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return str(int(v))
    return repr(float(v)) if isinstance(v, float) else str(v)


def _style(cell_xml: str) -> str:
    m = re.search(r'\bs="(\d+)"', cell_xml.split(">", 1)[0])
    return m.group(1) if m else ""


def _set_cell(xml: str, ref: str, value: Any, cached: Any = None, *, keep_formula_only_value: bool = False) -> str:
    m = re.fullmatch(r"([A-Z]+)(\d+)", ref)
    if not m:
        raise XlsxError(f"bad cell {ref}")
    col, row = fx.col_index(m.group(1)), int(m.group(2))
    sd = re.search(r"<sheetData\b[^>]*?(?:/>|>(.*?)</sheetData>)", xml, re.S)
    if sd is None:
        raise XlsxError("sheet has no data section")
    cm = re.search(_CELL.format(ref=ref), xml, re.S)
    if cm is not None:
        old = cm.group(0)
        if keep_formula_only_value:   # update only the cached value of an existing formula cell
            fm = re.search(r"<f\b[^>]*?(?:/>|>.*?</f>)", old, re.S)
            if fm is None:
                return xml
            head = re.sub(r'\s+t="[^"]*"', "", old.split(">", 1)[0].rstrip("/"))
            if cached is None or (isinstance(cached, str) and cached in fx.ERRORS):
                new = f"{head}>{fm.group(0)}</c>"
            elif isinstance(cached, bool):
                new = f'{head} t="b">{fm.group(0)}<v>{int(cached)}</v></c>'
            elif isinstance(cached, (int, float)):
                new = f"{head}>{fm.group(0)}<v>{_num(cached)}</v></c>"
            else:
                new = f'{head} t="str">{fm.group(0)}<v>{_xml_text(str(cached))}</v></c>'
        else:
            if re.search(r'<f[^>]*t="(?:shared|array)"[^>]*ref=', old):
                raise XlsxError(f"{ref} is the source of a shared or array formula; edit the cells it fills instead")
            new = _cell_xml(ref, value, _style(old), cached)
        return xml[:cm.start()] + new + xml[cm.end():]
    if keep_formula_only_value:
        return xml
    new_cell = _cell_xml(ref, value, "", cached)
    rm = re.search(_ROW.format(r=row), xml, re.S)
    if rm is not None:
        row_xml = rm.group(0)
        if row_xml.endswith("/>"):
            new_row = row_xml[:-2] + ">" + new_cell + "</row>"
        else:
            body_start = row_xml.index(">") + 1
            cells = list(re.finditer(r'<c\b[^>]*\br="([A-Z]+)\d+"[^>]*?(?:/>|>.*?</c>)', row_xml[body_start:], re.S))
            insert_at = len(row_xml) - len("</row>")
            for c in cells:
                if fx.col_index(c.group(1)) > col:
                    insert_at = body_start + c.start()
                    break
            new_row = row_xml[:insert_at] + new_cell + row_xml[insert_at:]
            new_row = re.sub(r'\s+spans="[^"]*"', "", new_row, count=1)   # spans are optional hints
        return xml[:rm.start()] + new_row + xml[rm.end():]
    # A new row, in row order.
    new_row = f'<row r="{row}">{new_cell}</row>'
    if sd.group(0).endswith("/>"):
        return xml[:sd.start()] + sd.group(0)[:-2] + ">" + new_row + "</sheetData>" + xml[sd.end():]
    body_start = sd.start(1)
    for r in re.finditer(r'<row\b[^>]*\br="(\d+)"', xml[body_start:sd.end(1)]):
        if int(r.group(1)) > row:
            pos = body_start + r.start()
            return xml[:pos] + new_row + xml[pos:]
    pos = sd.end(1)
    return xml[:pos] + new_row + xml[pos:]


def patch(book: Book, updates: dict[tuple[str, str], Any], out: Path) -> dict:
    """Write `book` with `updates` applied to `out` (a new file). Returns the recomputed values."""
    from jarvis.capabilities.documents import formula_allowed

    for (sheet, ref), value in updates.items():
        if sheet not in book.sheets:
            raise XlsxError(f"no sheet named {sheet!r}")
        if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d{0,6}", ref):
            raise XlsxError(f"bad cell reference {ref!r}")
        if _is_formula(value) and not formula_allowed(value):
            raise XlsxError(f"formula not allowed: {value}")
    values, unsupported = recompute(book, updates)
    zin = _check_zip(book.path)
    by_part: dict[str, list[tuple[str, Any, bool]]] = {}
    for (sheet, ref), value in updates.items():
        by_part.setdefault(book.sheets[sheet].xml_path, []).append((ref, value, False))
    for (sheet, ref), v in values.items():
        if (sheet, ref) not in updates and book.sheets[sheet].cached.get(ref) != v:
            by_part.setdefault(book.sheets[sheet].xml_path, []).append((ref, v, True))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename in by_part:
                xml = data.decode("utf-8")
                for ref, value, cached_only in by_part[info.filename]:
                    if cached_only:
                        xml = _set_cell(xml, ref, None, value, keep_formula_only_value=True)
                    else:
                        cached = values.get((_sheet_of(book, info.filename), ref)) if _is_formula(value) else None
                        xml = _set_cell(xml, ref, value, cached)
                data = xml.encode("utf-8")
            elif info.filename == "xl/workbook.xml":
                data = _full_calc_on_load(data.decode("utf-8")).encode("utf-8")
            elif info.filename == "xl/calcChain.xml":
                continue   # rebuilt by the app; a stale chain can make Excel "repair" the file
            zout.writestr(info, data)
    zin.close()
    payload = buf.getvalue()
    with zipfile.ZipFile(book.path) as original:
        had_chain = "xl/calcChain.xml" in original.namelist()
    if had_chain:
        payload = _drop_calc_chain_refs(payload)
    out.write_bytes(payload)
    return {"values": values, "unsupported": unsupported}


def _sheet_of(book: Book, part: str) -> str:
    return next((n for n, s in book.sheets.items() if s.xml_path == part), "")


def _full_calc_on_load(xml: str) -> str:
    if "<calcPr" in xml:
        if "fullCalcOnLoad" in xml:
            return re.sub(r'fullCalcOnLoad="[^"]*"', 'fullCalcOnLoad="1"', xml)
        return re.sub(r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', xml, count=1)
    # calcPr comes after sheets, functionGroups, externalReferences and definedNames (schema order).
    for tag in ("</definedNames>", "</externalReferences>", "</functionGroups>", "</sheets>"):
        if tag in xml:
            return xml.replace(tag, tag + '<calcPr fullCalcOnLoad="1"/>', 1)
    return xml


def _drop_calc_chain_refs(payload: bytes) -> bytes:
    """Remove the calcChain relationship and content-type entries after dropping the part itself."""
    zin = zipfile.ZipFile(io.BytesIO(payload))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/_rels/workbook.xml.rels":
                data = re.sub(rb"<Relationship\b[^>]*calcChain[^>]*/>", b"", data)
            elif info.filename == "[Content_Types].xml":
                data = re.sub(rb"<Override\b[^>]*calcChain[^>]*/>", b"", data)
            zout.writestr(info, data)
    return buf.getvalue()
