"""Layout checks for PowerPoint files without rendering them.

Finds shapes outside the slide, text boxes whose text cannot fit (estimated from font size, box
size and line wrapping), and text boxes that overlap. An estimate: it errs towards reporting, and
GenOffice's renderer-based audit replaces it when GenOffice is installed.
"""

from __future__ import annotations

import math
from pathlib import Path

EMU_PER_PT = 12700
DEFAULT_PT = 18.0


def _font_pt(frame) -> float:
    sizes = [r.font.size.pt for p in frame.paragraphs for r in p.runs if r.font.size is not None]
    return max(sizes) if sizes else DEFAULT_PT


def _lines(text: str, width_pt: float, font_pt: float) -> int:
    chars_per_line = max(1, int(width_pt / (font_pt * 0.5)))
    return sum(max(1, math.ceil(len(line) / chars_per_line)) for line in text.split("\n"))


def audit(path: Path) -> dict:
    from pptx import Presentation
    prs = Presentation(str(path))
    sw, sh = prs.slide_width, prs.slide_height
    issues = []
    for n, slide in enumerate(prs.slides, start=1):
        boxes = []
        for shape in slide.shapes:
            if shape.left is None or shape.width is None:
                continue
            left, t, w, h = shape.left, shape.top, shape.width, shape.height
            if left < 0 or t < 0 or left + w > sw * 1.001 or t + h > sh * 1.001:
                issues.append({"slide": n, "shape": shape.name, "type": "off_slide",
                               "message": f"'{shape.name}' extends beyond the slide"})
            if shape.has_text_frame and shape.text_frame.text.strip():
                from pptx.enum.text import MSO_AUTO_SIZE
                frame = shape.text_frame
                pt = _font_pt(frame)
                width_pt = max(1.0, w / EMU_PER_PT - 14)
                wraps = frame.word_wrap is not False
                lines = _lines(frame.text, width_pt, pt) if wraps else len(frame.text.splitlines() or [""])
                needed = lines * pt * 1.2 + 10
                auto = frame.auto_size
                if auto == MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT:
                    # The box grows with its text: a problem only if it then leaves the slide.
                    if t + needed * EMU_PER_PT > sh * 1.001:
                        issues.append({"slide": n, "shape": shape.name, "type": "off_slide",
                                       "message": f"'{shape.name}' grows past the bottom of the slide"})
                elif auto != MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE and needed > h / EMU_PER_PT * 1.05:
                    issues.append({"slide": n, "shape": shape.name, "type": "overflow",
                                   "message": f"text in '{shape.name}' needs ~{needed:.0f}pt but the box is "
                                              f"{h / EMU_PER_PT:.0f}pt tall"})
                boxes.append((shape.name, left, t, left + w, t + h))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                a, b = boxes[i], boxes[j]
                ix = max(0, min(a[3], b[3]) - max(a[1], b[1]))
                iy = max(0, min(a[4], b[4]) - max(a[2], b[2]))
                small = min((a[3] - a[1]) * (a[4] - a[2]), (b[3] - b[1]) * (b[4] - b[2])) or 1
                if ix * iy / small > 0.15:
                    issues.append({"slide": n, "shape": f"{a[0]} / {b[0]}", "type": "overlap",
                                   "message": f"'{a[0]}' and '{b[0]}' overlap"})
    return {"slides": len(prs.slides), "issues": issues, "method": "estimate (python-pptx geometry)"}


def slide_texts(path: Path) -> list[dict]:
    from pptx import Presentation
    prs = Presentation(str(path))
    out = []
    for n, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
        texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        out.append({"slide": n, "title": title, "text": "\n".join(texts)[:1500]})
    return out
