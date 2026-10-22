"""Reports from verified numbers: a before/after comparison as PDF, Word or PowerPoint.

Every figure in a report comes from the caller (normally a workbook edit Jarvis just made and
recalculated), never from the model's imagination, and the report states its source and time.
Charts are real: drawn as vectors in the PDF, a native (editable) chart in PowerPoint, and an image
in Word (python-docx cannot write native charts).
"""

from __future__ import annotations

import io
import time
from dataclasses import dataclass

ACCENT = (79, 157, 255)
BEFORE = (148, 163, 184)
UP = (220, 38, 38)
DOWN = (22, 163, 74)
INK = (30, 32, 38)
MUTED = (100, 104, 114)


@dataclass
class Item:
    label: str
    before: float
    after: float

    @property
    def delta(self) -> float:
        return self.after - self.before

    @property
    def pct(self) -> float | None:
        return (self.delta / self.before * 100) if self.before else None


def items_from(raw: list[dict]) -> list[Item]:
    out = []
    for r in raw:
        try:
            out.append(Item(str(r["label"])[:80], float(r["before"]), float(r["after"])))
        except (KeyError, TypeError, ValueError):
            continue
    return out


def fmt(x: float, unit: str = "") -> str:
    s = f"{x:,.2f}".rstrip("0").rstrip(".").replace(",", " ")
    return f"{s} {unit}".strip()


def fmt_delta(it: Item, unit: str = "") -> str:
    sign = "+" if it.delta > 0 else "−" if it.delta < 0 else "±"
    pct = f" ({sign}{abs(it.pct):.1f}%)" if it.pct is not None else ""
    return f"{sign}{fmt(abs(it.delta), unit)}{pct}"


L = {"en": {"item": "Item", "before": "Before", "after": "After", "change": "Change", "source": "Source",
            "made": "Prepared by Jarvis", "notes": "Notes", "total": "Total"},
     "ru": {"item": "Статья", "before": "Было", "after": "Стало", "change": "Изменение", "source": "Источник",
            "made": "Подготовлено Jarvis", "notes": "Примечания", "total": "Итого"},
     "pt": {"item": "Item", "before": "Antes", "after": "Depois", "change": "Variação", "source": "Fonte",
            "made": "Preparado pelo Jarvis", "notes": "Notas", "total": "Total"}}


def _labels(lang: str, before_label: str, after_label: str) -> dict:
    lab = dict(L.get(lang, L["en"]))
    if before_label:
        lab["before"] = before_label
    if after_label:
        lab["after"] = after_label
    return lab


def pdf(title: str, items: list[Item], *, source: str = "", notes: list[str] | None = None, unit: str = "",
        lang: str = "en", before_label: str = "", after_label: str = "", total: Item | None = None) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    from jarvis.capabilities.documents import _pdf_fonts
    regular, bold = _pdf_fonts()
    lab = _labels(lang, before_label, after_label)
    pdf = FPDF(format="A4")
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, margin=16)
    pdf.add_font("body", "", regular)
    pdf.add_font("body", "B", bold or regular)
    pdf.set_title(title[:200])
    pdf.set_creator("Jarvis WorkOS")
    pdf.add_page()
    nl = dict(new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("body", "B", 19)
    pdf.set_text_color(*INK)
    pdf.multi_cell(0, 9, title, **nl)
    pdf.set_draw_color(*ACCENT)
    pdf.set_line_width(0.8)
    y = pdf.get_y() + 1.5
    pdf.line(pdf.l_margin, y, pdf.l_margin + 26, y)
    pdf.ln(5)
    pdf.set_font("body", "", 9.5)
    pdf.set_text_color(*MUTED)
    meta = f"{lab['made']} · {time.strftime('%Y-%m-%d %H:%M')}" + (f" · {lab['source']}: {source}" if source else "")
    pdf.multi_cell(0, 5, meta, **nl)
    pdf.ln(4)

    # Chart: paired horizontal bars, before (grey) and after (accent).
    rows = items[:14]
    if rows:
        top = max(max(abs(i.before), abs(i.after)) for i in rows) or 1
        label_w, bar_w, row_h = 52, pdf.epw - 52 - 30, 9
        pdf.set_font("body", "", 8.5)
        for it in rows:
            y0 = pdf.get_y()
            pdf.set_text_color(*INK)
            pdf.set_xy(pdf.l_margin, y0 + 1)
            pdf.cell(label_w - 2, 6, it.label[:34])
            x0 = pdf.l_margin + label_w
            pdf.set_fill_color(*BEFORE)
            pdf.rect(x0, y0 + 1, max(0.4, bar_w * abs(it.before) / top), 3, style="F")
            pdf.set_fill_color(*ACCENT)
            pdf.rect(x0, y0 + 4.6, max(0.4, bar_w * abs(it.after) / top), 3, style="F")
            pdf.set_text_color(*(UP if it.delta > 0 else DOWN if it.delta < 0 else MUTED))
            pdf.set_xy(x0 + bar_w + 2, y0 + 1)
            pdf.cell(28, 6, f"{'+' if it.delta > 0 else ''}{it.pct:.1f}%" if it.pct is not None else "")
            pdf.set_y(y0 + row_h)
        pdf.set_font("body", "", 8)
        pdf.set_text_color(*MUTED)
        x = pdf.l_margin + label_w
        pdf.set_fill_color(*BEFORE)
        pdf.rect(x, pdf.get_y() + 1.5, 4, 3, style="F")
        pdf.set_xy(x + 5, pdf.get_y())
        pdf.cell(30, 6, lab["before"])
        pdf.set_fill_color(*ACCENT)
        pdf.rect(x + 36, pdf.get_y() + 1.5, 4, 3, style="F")
        pdf.set_xy(x + 41, pdf.get_y())
        pdf.cell(30, 6, lab["after"], **nl)
        pdf.ln(4)

    # Table.
    widths = [pdf.epw * f for f in (0.38, 0.19, 0.19, 0.24)]
    pdf.set_font("body", "B", 9.5)
    pdf.set_text_color(*INK)
    pdf.set_fill_color(241, 244, 249)
    for w, h in zip(widths, (lab["item"], lab["before"], lab["after"], lab["change"])):
        pdf.cell(w, 8, h, border="B", fill=True, align="L" if h == lab["item"] else "R")
    pdf.ln(8)
    pdf.set_font("body", "", 9.5)
    for it in items + ([total] if total else []):
        bold_row = total is not None and it is total
        pdf.set_font("body", "B" if bold_row else "", 9.5)
        pdf.set_text_color(*INK)
        pdf.cell(widths[0], 7, (lab["total"] if bold_row and not it.label else it.label)[:48], border="B")
        pdf.cell(widths[1], 7, fmt(it.before, unit), border="B", align="R")
        pdf.cell(widths[2], 7, fmt(it.after, unit), border="B", align="R")
        pdf.set_text_color(*(UP if it.delta > 0 else DOWN if it.delta < 0 else MUTED))
        pdf.cell(widths[3], 7, fmt_delta(it, unit), border="B", align="R")
        pdf.ln(7)
    if notes:
        pdf.ln(5)
        pdf.set_font("body", "B", 11)
        pdf.set_text_color(*INK)
        pdf.cell(0, 7, lab["notes"], **nl)
        pdf.set_font("body", "", 9.5)
        pdf.set_text_color(*MUTED)
        for n in notes[:20]:
            pdf.multi_cell(0, 5.2, f"• {n}", **nl)
    out = pdf.output()
    return bytes(out)


def chart_png(items: list[Item], *, lang: str = "en", before_label: str = "", after_label: str = "",
              width: int = 1400) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    from jarvis.capabilities.documents import find_fonts
    lab = _labels(lang, before_label, after_label)
    rows = items[:14]
    row_h, pad, label_w = 64, 40, 420
    height = pad * 2 + row_h * len(rows) + 60
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    fonts = find_fonts()
    try:
        font = ImageFont.truetype(fonts[0], 26) if fonts else ImageFont.load_default()
        small = ImageFont.truetype(fonts[0], 22) if fonts else ImageFont.load_default()
    except OSError:
        font = small = ImageFont.load_default()
    top = max((max(abs(i.before), abs(i.after)) for i in rows), default=1) or 1
    bar_w = width - label_w - pad * 2 - 140
    for n, it in enumerate(rows):
        y = pad + n * row_h
        d.text((pad, y + 14), it.label[:30], fill=INK, font=font)
        x = pad + label_w
        d.rectangle([x, y + 10, x + max(2, int(bar_w * abs(it.before) / top)), y + 28], fill=BEFORE)
        d.rectangle([x, y + 34, x + max(2, int(bar_w * abs(it.after) / top)), y + 52], fill=ACCENT)
        if it.pct is not None:
            d.text((x + bar_w + 16, y + 18), f"{'+' if it.delta > 0 else ''}{it.pct:.1f}%",
                   fill=UP if it.delta > 0 else DOWN if it.delta < 0 else MUTED, font=small)
    y = pad + len(rows) * row_h + 14
    x = pad + label_w
    d.rectangle([x, y + 6, x + 22, y + 24], fill=BEFORE)
    d.text((x + 30, y), lab["before"], fill=MUTED, font=small)
    d.rectangle([x + 200, y + 6, x + 222, y + 24], fill=ACCENT)
    d.text((x + 230, y), lab["after"], fill=MUTED, font=small)
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    return buf.getvalue()


def docx(title: str, items: list[Item], *, source: str = "", notes: list[str] | None = None, unit: str = "",
         lang: str = "en", before_label: str = "", after_label: str = "", total: Item | None = None) -> bytes:
    from docx import Document
    from docx.shared import Cm, Pt

    lab = _labels(lang, before_label, after_label)
    doc = Document()
    doc.add_heading(title, level=0)
    meta = doc.add_paragraph(f"{lab['made']} · {time.strftime('%Y-%m-%d %H:%M')}"
                             + (f" · {lab['source']}: {source}" if source else ""))
    meta.runs[0].font.size = Pt(9)
    if items:
        doc.add_picture(io.BytesIO(chart_png(items, lang=lang, before_label=before_label, after_label=after_label)),
                        width=Cm(16))
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light List Accent 1" if "Light List Accent 1" in [s.name for s in doc.styles] else table.style
    for cell, h in zip(table.rows[0].cells, (lab["item"], lab["before"], lab["after"], lab["change"])):
        cell.text = h
    for it in items + ([total] if total else []):
        cells = table.add_row().cells
        cells[0].text = it.label or lab["total"]
        cells[1].text = fmt(it.before, unit)
        cells[2].text = fmt(it.after, unit)
        cells[3].text = fmt_delta(it, unit)
    if notes:
        doc.add_heading(lab["notes"], level=2)
        for n in notes[:20]:
            doc.add_paragraph(n, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def pptx(title: str, items: list[Item], *, source: str = "", notes: list[str] | None = None, unit: str = "",
         lang: str = "en", before_label: str = "", after_label: str = "", total: Item | None = None,
         points: list[str] | None = None) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Emu, Pt

    lab = _labels(lang, before_label, after_label)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 16:9
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.placeholders[1].text = f"{lab['made']} · {time.strftime('%Y-%m-%d')}" + (f"\n{lab['source']}: {source}" if source else "")
    if points:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = lab["notes"] if not notes else title
        body = s.placeholders[1].text_frame
        body.text = points[0]
        for p in points[1:8]:
            body.add_paragraph().text = p
    if items:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = f"{lab['before']} → {lab['after']}"
        data = CategoryChartData()
        data.categories = [i.label[:30] for i in items[:12]]
        data.add_series(lab["before"], [i.before for i in items[:12]])
        data.add_series(lab["after"], [i.after for i in items[:12]])
        frame = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Emu(600000), Emu(1500000), Emu(10992000),
                                   Emu(4900000), data)
        chart = frame.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        plot = chart.plots[0]
        plot.series[0].format.fill.solid()
        plot.series[0].format.fill.fore_color.rgb = RGBColor(*BEFORE)
        plot.series[1].format.fill.solid()
        plot.series[1].format.fill.fore_color.rgb = RGBColor(*ACCENT)
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = lab["change"]
        rows = items[:12] + ([total] if total else [])
        shape = s.shapes.add_table(len(rows) + 1, 4, Emu(600000), Emu(1500000), Emu(10992000),
                                   Emu(min(4900000, 380000 * (len(rows) + 1))))
        t = shape.table
        for c, h in enumerate((lab["item"], lab["before"], lab["after"], lab["change"])):
            t.cell(0, c).text = h
        for r, it in enumerate(rows, start=1):
            for c, v in enumerate((it.label or lab["total"], fmt(it.before, unit), fmt(it.after, unit), fmt_delta(it, unit))):
                t.cell(r, c).text = v
                for p in t.cell(r, c).text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(14)
    if notes:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = lab["notes"]
        body = s.placeholders[1].text_frame
        body.text = notes[0]
        for n in notes[1:8]:
            body.add_paragraph().text = n
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
