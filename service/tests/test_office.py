"""Office files (Tests E–H): real .docx/.xlsx/.pptx/.pdf files written, edited, re-opened and checked;
approvals before edits, version history and undo, formula recalculation that keeps charts intact,
reports built only from verified numbers, and the file-safety policy."""

from __future__ import annotations

import asyncio
import zipfile
from pathlib import Path

import pytest

from jarvis import config as configmod
from jarvis.office import formulas as fx
from jarvis.office import paths, pptx_check, report, xlsx
from jarvis.tools import CONFIRM, SAFE

docx_mod = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")
pptx_mod = pytest.importorskip("pptx")
pypdf = pytest.importorskip("pypdf")


def run(coro):
    return asyncio.run(coro)


class FakeBrain:
    gemini_key = None
    claude_key = "x"
    provider = "claude"
    failures: dict = {}


@pytest.fixture
def jarvis(tmp_path, monkeypatch, fake):
    monkeypatch.setenv("JARVIS_NO_VOICE", "1")
    monkeypatch.setattr(configmod, "data_dir", lambda: tmp_path / "data")
    monkeypatch.setattr(configmod, "config_dir", lambda: tmp_path / "cfg")
    monkeypatch.setenv("XDG_CACHE_HOME", str(tmp_path / "cache"))
    from jarvis import daemon
    cfg = configmod.Config(path=tmp_path / "cfg" / "config.toml")
    cfg.values["files"]["roots"] = [str(tmp_path)]
    cfg.values["documents"]["dir"] = str(tmp_path / "docs")
    j = daemon.Jarvis(cfg)
    j.platform = j.ctx.platform = fake
    j.brain = j.ctx.brain = FakeBrain()
    return j


async def call(j, name, args, allow=None, wait=20.0):
    holder, seen = {}, []

    async def work(task):
        holder["r"] = await j.run_tool(task, name, args)
        return "x"
    task = j.tasks.start(name, "text", work)
    for _ in range(int(wait / 0.02)):
        if task.handle.done():
            break
        for a in list(j.tasks.approvals.values()):
            if a.task_id == task.id and a not in seen:
                seen.append(a)
                if allow is not None:
                    j.tasks.resolve(a.id, allow)
        await asyncio.sleep(0.02)
    if not task.handle.done():
        task.handle.cancel()
    try:
        await task.handle
    except BaseException:
        pass
    return holder.get("r") or {"ok": False, "message": task.error or task.state.value}, seen


def budget(path: Path) -> Path:
    from openpyxl.chart import BarChart, Reference
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    rows = [("Category", "Estimate", "Actual"), ("Construction", 120000, 118500), ("Flooring", 18000, 0),
            ("Electrical", 9500, 9100), ("Kitchen", 22000, 0), ("Total", "=SUM(B2:B5)", "=SUM(C2:C5)")]
    for r in rows:
        ws.append(r)
    ws["E1"] = "Contingency 10%"
    ws["E2"] = "=ROUND(B6*0.1,2)"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    ws.add_chart(chart, "G2")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# ------------------------------------------------------------------ Test F: spreadsheets

def test_f_scale_asks_first_recalculates_keeps_chart_and_undo_restores(jarvis, tmp_path):
    book = budget(tmp_path / "Projects" / "Project budget.xlsx")
    original = book.read_bytes()

    async def main():
        found, _ = await call(jarvis, "office", {"action": "latest", "query": "budget", "kind": "xlsx"})
        assert found["ok"] and found["path"] == str(book)
        r, seen = await call(jarvis, "office", {"action": "scale", "path": str(book), "match": "construction",
                                                "percent": 8, "column": "Estimate"}, allow=False)
        assert seen and seen[0].risk == CONFIRM and "120 000 → 129 600" in seen[0].summary
        assert book.read_bytes() == original                         # denied: untouched
        r, seen = await call(jarvis, "office", {"action": "scale", "path": str(book), "match": "construction",
                                                "percent": 8, "column": "Estimate"}, allow=True)
        assert r["ok"] and r["verified"], r
        changes = {c["label"]: (c["before"], c["after"]) for c in r["changes"]}
        assert changes["Construction"] == (120000, 129600)
        assert changes["Total"] == (169500, 179100)                  # the formula total, recalculated
        assert any(c["formula"] == "=ROUND(B6*0.1,2)" for c in r["changes"] if c.get("kind") == "formula")
        # Re-open the real file: value written, formulas kept, cached totals updated, chart still there.
        wb = openpyxl.load_workbook(book)
        assert wb["Budget"]["B2"].value == 129600 and wb["Budget"]["B6"].value == "=SUM(B2:B5)"
        assert openpyxl.load_workbook(book, data_only=True)["Budget"]["B6"].value == 179100
        with zipfile.ZipFile(book) as z:
            assert "xl/charts/chart1.xml" in z.namelist()
            assert 'fullCalcOnLoad="1"' in z.read("xl/workbook.xml").decode()
        card = jarvis.hub.cards[-1]
        assert card["kind"] == "comparison" and card["provenance"]["kind"] == "calculated"
        # "Show me what changed" against the saved version.
        cmp, _ = await call(jarvis, "office", {"action": "compare", "path": str(book)})
        assert {c["label"] for c in cmp["changes"]} >= {"Construction", "Total"}
        r, _ = await call(jarvis, "undo", {}, allow=True)
        assert r["ok"] and openpyxl.load_workbook(book)["Budget"]["B2"].value == 120000
    run(main())


def test_f_ambiguous_rows_and_columns_ask_instead_of_guessing(jarvis, tmp_path):
    book = budget(tmp_path / "b.xlsx")

    async def main():
        r, _ = await call(jarvis, "office", {"action": "scale", "path": str(book), "match": "construction", "percent": 5,
                                             "dry_run": True})
        assert not r["ok"] and r["needs_choice"] and any("Actual" in c for c in r["choices"])
        r, seen = await call(jarvis, "office", {"action": "scale", "path": str(book), "match": "construction",
                                                "percent": 5, "column": "B", "dry_run": True})
        assert r["ok"] and not seen and r["plan"][0]["after"] == 126000       # dry run: no approval, no change
        r, _ = await call(jarvis, "office", {"action": "scale", "path": str(book), "match": "plumbing", "percent": 5,
                                             "dry_run": True})
        assert not r["ok"] and "No row" in r["message"]
    run(main())


def test_f_update_cells_and_formula_whitelist(jarvis, tmp_path):
    book = budget(tmp_path / "b.xlsx")

    async def main():
        r, seen = await call(jarvis, "office", {"action": "update", "path": str(book), "changes": [
            {"cell": "B3", "value": "20 160"}, {"cell": "D6", "value": "=B6-C6"}]}, allow=True)
        assert r["ok"] and seen
        wb = openpyxl.load_workbook(book, data_only=True)["Budget"]
        assert wb["B3"].value == 20160 and wb["D6"].value == (120000 + 20160 + 9500 + 22000) - (118500 + 9100)
        r, _ = await call(jarvis, "office", {"action": "update", "path": str(book), "changes": [
            {"cell": "A9", "value": '=WEBSERVICE("http://evil")'}]}, allow=True)
        assert not r["ok"] and "not allowed" in r["message"]
    run(main())


def test_formula_engine_matches_spreadsheet_semantics():
    wb = fx.Workbook({"S": {"A1": 10, "A2": "x", "A3": 5, "B1": "=SUM(A1:A3)", "B2": "=AVERAGE(A1:A3)", "B3": "=A1/0",
                            "B4": "=IF(B1>10,\"high\",\"low\")", "B5": "=B5+1", "B6": "=ROUND(2.345,2)",
                            "B7": "=IFERROR(A1/0,-1)", "B8": "=MAX(A1,A3)*-1"}})
    assert [wb.value("S", c) for c in ("B1", "B2", "B3", "B4", "B5", "B6", "B7", "B8")] == \
        [15, 7.5, "#DIV/0!", "high", "#CIRC!", 2.35, -1, -10]
    with pytest.raises(fx.Unsupported):
        fx.Workbook({"S": {"A1": "=VLOOKUP(1,B1:C2,2)"}}).value("S", "A1")


def test_xlsx_check_reports_formula_errors(jarvis, tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["A2"] = 0, "=10/A1"
    wb.save(tmp_path / "e.xlsx")
    r, _ = run(call(jarvis, "office", {"action": "check", "path": str(tmp_path / "e.xlsx")}))
    assert r["ok"] and r["errors"][0]["error"] == "#DIV/0!"


# ------------------------------------------------------------------ Test E: Word

def test_e_docx_section_edit_persists_and_undo(jarvis, tmp_path):
    doc = docx_mod.Document()
    doc.add_heading("Kitchen project", 0)
    doc.add_heading("Scope", 1)
    doc.add_paragraph("Old scope paragraph one.")
    doc.add_paragraph("Old scope paragraph two.")
    doc.add_heading("Schedule", 1)
    doc.add_paragraph("Start in November.")
    path = tmp_path / "Kitchen.docx"
    doc.save(path)

    async def main():
        r, seen = await call(jarvis, "office", {"action": "replace_section", "path": str(path), "heading": "scope",
                                                "paragraphs": ["New cabinets in oak.", "Countertop: quartz, 20 mm."]},
                             allow=True)
        assert r["ok"] and seen and "New cabinets in oak." in seen[0].summary
        texts = [p.text for p in docx_mod.Document(path).paragraphs]
        assert texts == ["Kitchen project", "Scope", "New cabinets in oak.", "Countertop: quartz, 20 mm.", "Schedule",
                         "Start in November."]
        r, _ = await call(jarvis, "office", {"action": "replace_text", "path": str(path), "find": "November",
                                             "replace": "December"}, allow=True)
        assert r["ok"] and "Start in December." in [p.text for p in docx_mod.Document(path).paragraphs]
        cmp, _ = await call(jarvis, "office", {"action": "compare", "path": str(path)})
        assert cmp["changes"] and "December" in cmp["changes"][0]["after"]
        r, _ = await call(jarvis, "undo", {}, allow=True)
        assert "Start in November." in [p.text for p in docx_mod.Document(path).paragraphs]
    run(main())


# ------------------------------------------------------------------ Tests G and H: PPTX and PDF reports

def test_g_h_reports_from_the_verified_change(jarvis, tmp_path):
    book = budget(tmp_path / "budget.xlsx")

    async def main():
        r, _ = await call(jarvis, "office", {"action": "scale", "path": str(book), "match": "flooring", "percent": 12,
                                             "column": "Estimate"}, allow=True)
        assert r["ok"]
        pdf, seen = await call(jarvis, "office", {"action": "report", "format": "pdf", "title": "Flooring +12%"})
        assert pdf["ok"] and not seen
        text = "".join(p.extract_text() for p in pypdf.PdfReader(pdf["path"]).pages)
        assert "Flooring" in text and "20 160" in text and "18 000" in text and "budget.xlsx" in text
        deck, _ = await call(jarvis, "office", {"action": "report", "format": "pptx", "title": "Flooring +12%",
                                                "points": ["Flooring estimate raised by 12%"]})
        prs = pptx_mod.Presentation(deck["path"])
        charts = [s for slide in prs.slides for s in slide.shapes if s.has_chart]
        assert len(prs.slides) >= 3 and charts
        assert list(charts[0].chart.plots[0].series[1].values)[0] == 20160      # the real, editable chart
        docx_r, _ = await call(jarvis, "office", {"action": "report", "format": "docx", "title": "Flooring +12%"})
        assert docx_r["ok"] and Path(docx_r["path"]).stat().st_size > 5000
        audit = pptx_check.audit(Path(deck["path"]))
        assert audit["slides"] == len(prs.slides)
    run(main())


def test_report_without_a_verified_change_refuses(jarvis):
    from jarvis.capabilities import office
    office.LAST_CHANGE.clear()
    r, _ = run(call(jarvis, "office", {"action": "report", "format": "pdf"}))
    assert not r["ok"] and "no verified change" in r["message"].lower()


def test_g_layout_check_finds_overflow_and_off_slide(tmp_path):
    from pptx.util import Emu, Pt
    prs = pptx_mod.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(100000), Emu(100000), Emu(2000000), Emu(400000))
    from pptx.enum.text import MSO_AUTO_SIZE
    box.text_frame.word_wrap = True
    box.text_frame.auto_size = MSO_AUTO_SIZE.NONE      # a fixed box, as in most templates
    box.text_frame.text = "A very long sentence that cannot possibly fit into this small box " * 6
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    slide.shapes.add_textbox(prs.slide_width - Emu(100000), Emu(100000), Emu(2000000), Emu(400000)).text_frame.text = "x"
    prs.save(tmp_path / "bad.pptx")
    kinds = {i["type"] for i in pptx_check.audit(tmp_path / "bad.pptx")["issues"]}
    assert {"overflow", "off_slide"} <= kinds


def test_report_numbers_and_signs():
    it = report.Item("Flooring", 18000, 20160)
    assert report.fmt_delta(it) == "+2 160 (+12.0%)" and report.fmt(1234567.5) == "1 234 567.5"


# ------------------------------------------------------------------ file safety

def test_paths_refuse_traversal_secrets_and_outside_roots(jarvis, tmp_path):
    cfg = jarvis.config
    (tmp_path / ".ssh").mkdir()
    (tmp_path / ".ssh" / "config").write_text("x")
    (tmp_path / "ok.txt").write_text("x")
    assert paths.check_readable(cfg, str(tmp_path / "ok.txt")) == (tmp_path / "ok.txt").resolve()
    for bad in (str(tmp_path / ".ssh" / "config"), str(tmp_path / "docs" / ".." / ".." / "etc" / "passwd"),
                str(Path.home() / "outside-test.txt"), "file://" + str(tmp_path / ".ssh" / "config")):
        with pytest.raises((paths.PathRefused, FileNotFoundError)):
            paths.check_readable(cfg, bad)
    (configmod.data_dir()).mkdir(parents=True, exist_ok=True)
    (configmod.data_dir() / "memory.db").write_text("x")
    with pytest.raises(paths.PathRefused):
        paths.check_readable(cfg, str(configmod.data_dir() / "memory.db"))
    r, _ = run(call(jarvis, "office", {"action": "inspect", "path": str(tmp_path / ".ssh" / "config")}))
    assert not r["ok"] and "credentials" in r["message"]


def test_office_risk_levels():
    from jarvis.tools import REGISTRY
    t = REGISTRY["office"]
    assert t.risk_for({"action": "scale"}) == CONFIRM and t.risk_for({"action": "scale", "dry_run": True}) == SAFE
    for a in ("latest", "inspect", "read", "compare", "report", "preview", "check", "convert"):
        assert t.risk_for({"action": a}) == SAFE
    for a in ("update", "replace_text", "replace_section", "append", "restore"):
        assert t.risk_for({"action": a}) == CONFIRM


def test_xlsx_patch_inserts_new_rows_and_cells_in_order(tmp_path):
    p = budget(tmp_path / "x.xlsx")
    book = xlsx.load(p)
    out = tmp_path / "y.xlsx"
    xlsx.patch(book, {("Budget", "A9"): "Notes", ("Budget", "C9"): 1, ("Budget", "B9"): "mid", ("Budget", "A8"): "eight"}, out)
    ws = openpyxl.load_workbook(out)["Budget"]
    assert [ws["A8"].value, ws["A9"].value, ws["B9"].value, ws["C9"].value] == ["eight", "Notes", "mid", 1]
